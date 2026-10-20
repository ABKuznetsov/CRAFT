"""Editable Word publication of the same frozen content used by the preview."""

from io import BytesIO
import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from craft.analysis.reporting.automatic import composed_figures, publication_tables, report_paragraphs


def _property(parent, tag, **attributes):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    parent.append(element)
    return element


def _document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    for name in ("Normal", "Title", "Subtitle", "Caption", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + key), None)
        for tag in ("w:pBdr", "w:numPr"):
            element = style.element.find(".//" + qn(tag))
            if element is not None:
                element.getparent().remove(element)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    _property(normal.element.get_or_add_rPr(), "w:lang", val="en-US")
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Caption"].font.size = Pt(10)
    doc.styles["Caption"].font.bold = False
    doc.add_paragraph("Crystal structure report", "Title")
    doc.add_paragraph(title, "Subtitle")
    doc.core_properties.title = "Crystal structure report " + title
    doc.core_properties.language = "en-US"
    doc.core_properties.author = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _property(footer._p, "w:fldSimple", instr="PAGE")
    return doc


def _note(doc, text):
    paragraph = doc.add_paragraph(text)
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("555555")


def _picture(paragraph, panel, width, height):
    try:
        image = Image.from_blob(panel.preview_png)
    except Exception as error:
        raise ValueError(f"The saved image for figure '{panel.title}' cannot be read.") from error
    scale = min(width / image.px_width, height / image.px_height)
    paragraph.add_run().add_picture(BytesIO(panel.preview_png),
        width=int(image.px_width * scale), height=int(image.px_height * scale))


def _figures(doc, draft, width):
    for group in composed_figures(draft):
        panels = group.panels
        if len(panels) == 1:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            _picture(paragraph, panels[0], width, Inches(4.8))
        else:
            table = doc.add_table(rows=(len(panels) + 1) // 2, cols=2)
            table.autofit = False
            for column in table.columns:
                column.width = width // 2
            for index, panel in enumerate(panels):
                cell = table.cell(index // 2, index % 2)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                _picture(paragraph, panel, width // 2 - Mm(4), Inches(3.0))
                label = cell.add_paragraph(f"({chr(ord('a') + index)})")
                label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for row in table.rows:
                _property(row._tr.get_or_add_trPr(), "w:cantSplit")
            for cell in table.rows[-1].cells:
                cell.paragraphs[-1].paragraph_format.keep_with_next = True
        caption = doc.add_paragraph(style="Caption")
        caption.add_run(f"Figure {group.number}. ").bold = True
        caption.add_run(panels[0].caption)
        if len(panels) > 1:
            for index, panel in enumerate(panels):
                if panel.caption != panels[0].caption:
                    caption.add_run(f" ({chr(ord('a') + index)}) {panel.caption}")
        if any(panel.stale for panel in panels):
            _note(doc, "This view belongs to an earlier analysis state.")


def _table(doc, source, width):
    columns = tuple(column for column in source.columns if column.visible)
    if not columns:
        _note(doc, "No columns are selected for publication.")
        return
    rows = tuple(row for row in source.rows if row.include_in_publication)
    table = doc.add_table(rows=1, cols=len(columns))
    table.autofit = False
    weights = [1.5 if c.id == "parameter" else 1. if c.id == "value" else
               1.2 if c.id == "label" else .85 if c.id in ("species", "tensor") else 1.
               for c in columns]
    widths = [int(width * weight / sum(weights)) for weight in weights]
    for column, size in zip(table.columns, widths):
        column.width = size
    borders = _property(table._tbl.tblPr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _property(borders, "w:" + side, val="single", sz=4, color="D9D9D9")
    _property(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")
    texts = [tuple(c.title + (f" ({c.unit})" if c.unit else "") for c in columns)]
    texts.extend(tuple(row.cells[c.id].display if c.id in row.cells else "—" for c in columns) for row in rows)
    for index, values in enumerate(texts):
        row = table.rows[0] if index == 0 else table.add_row()
        if max(map(len, values), default=0) < 1500:
            _property(row._tr.get_or_add_trPr(), "w:cantSplit")
        for cell, text, column, size in zip(row.cells, values, columns, widths):
            cell.width = size
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = _property(cell._tc.get_or_add_tcPr(), "w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                _property(margins, "w:" + side, w=65, type="dxa")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.keep_with_next = index == 0 or index == len(texts) - 2
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if column.id in
                                   ("parameter", "value", "label", "species") else WD_ALIGN_PARAGRAPH.CENTER)
            run = paragraph.add_run(text)
            run.font.size = Pt(9 if len(columns) > 6 else 10)
            run.bold = index == 0
            if index == 0:
                _property(cell._tc.get_or_add_tcPr(), "w:shd", fill="E8EEF3")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def export_report_docx(draft, report, path):
    """Write selected content atomically; figures remain embedded, tables editable."""
    doc = _document(report.structure_name)
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    for text in report_paragraphs(draft):
        doc.add_paragraph(text)
    _figures(doc, draft, width)
    for number, title, table, reason in publication_tables(draft, report):
        caption = doc.add_paragraph(style="Caption")
        caption.paragraph_format.keep_with_next = True
        caption.add_run(f"Table {number}. ").bold = True
        caption.add_run(title)
        if table is None:
            _note(doc, reason)
            continue
        _table(doc, table, width)
        if table.method:
            _note(doc, table.method)
        for warning in table.warnings:
            _note(doc, warning.message)
    target = Path(path)
    temporary = None
    try:
        with NamedTemporaryFile(dir=target.parent, suffix=".docx", delete=False) as stream:
            temporary = Path(stream.name)
        doc.save(temporary)
        os.replace(temporary, target)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
